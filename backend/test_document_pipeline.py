import sys
import os
import io
import shutil
import tempfile
import docx
import fitz  # PyMuPDF

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

sys.path.insert(0, r"c:\Users\umran\Desktop\RapidDoc_SGP")

from RapidDoc.backend.app.services.docx_editor import (
    get_docx_content,
    update_docx_content,
    find_replace_docx,
    find_text_variants,
    selective_replace_docx,
    apply_docx_styling,
    get_docx_images_count,
    get_docx_headers_footers
)
from RapidDoc.backend.app.services.pdf_editor import (
    get_pdf_content,
    update_pdf_content,
    find_replace_pdf,
    find_text_variants_pdf,
    selective_replace_pdf,
    apply_pdf_styling,
    get_pdf_images_count,
    get_pdf_headers_footers
)
from RapidDoc.backend.app.services.gemini_service import (
    understand_command,
    rewrite_text
)
from RapidDoc.backend.app.services.local_models import (
    classify_intent,
    rewrite_text_locally
)
from RapidDoc.backend.app.services.slot_extractor import extract_slots

def print_section(title):
    print("\n" + "=" * 65)
    print(f"📌 {title}")
    print("=" * 65)

def test_docx_pipeline():
    print_section("1. Testing DOCX Pipeline & Deterministic Handlers")
    
    tmp_dir = tempfile.mkdtemp(prefix="test_docx_")
    test_docx_path = os.path.join(tmp_dir, "sample_test.docx")
    
    doc = docx.Document()
    doc.add_heading("RapidDoc Project Assessment", level=1)
    doc.add_paragraph("This is the initial draft version of the system documentation.")
    doc.add_paragraph("Machine learning models are integrated into the dual brain framework.")
    doc.add_paragraph("Third paragraph for testing paragraph indexing.")
    doc.save(test_docx_path)
    print(f"Created temporary test document: {test_docx_path}")
    
    # Test Content Extraction
    content = get_docx_content(test_docx_path)
    print(f"Extracted {len(content)} paragraphs:")
    for item in content:
        print(f"  [{item['index']}]: {item['text']}")
    assert len(content) >= 3, "Failed to extract DOCX content"
    
    # Test Header & Footer Update
    success = apply_docx_styling(
        test_docx_path, test_docx_path,
        font_name="Arial",
        font_size=12.0,
        header_text="RapidDoc Confidential Header",
        footer_text="Page 1 - All Rights Reserved"
    )
    print(f"Apply styling & Header/Footer result: {success}")
    assert success, "Failed to apply DOCX styling & header/footer"
    
    hf_data = get_docx_headers_footers(test_docx_path)
    print(f"Verified Headers/Footers: {hf_data}")
    assert "RapidDoc Confidential Header" in hf_data["headers"]
    assert "Page 1 - All Rights Reserved" in hf_data["footers"]
    
    # Test Find Variants & Selective Replace
    variants = find_text_variants(test_docx_path, "draft", case_sensitive=False)
    print(f"Find variants for 'draft': {variants}")
    assert variants["total_matches"] >= 1, "Should find 'draft'"
    
    rep_res = selective_replace_docx(
        test_docx_path, test_docx_path,
        find_text="draft",
        replace_text="Production Release",
        selected_variants=[g["variant"] for g in variants["groups"]]
    )
    print(f"Selective replace result: {rep_res}")
    assert rep_res["matches_replaced"] >= 1
    
    # Test Content Block Update
    update_res = update_docx_content(
        test_docx_path, test_docx_path,
        [{"index": 2, "text": "Dual brain system with DistilBERT and T5 Seq2Seq is operational."}]
    )
    print(f"Update DOCX paragraph content: {update_res}")
    assert update_res, "Failed to update DOCX paragraph"
    
    shutil.rmtree(tmp_dir, ignore_errors=True)
    print("✅ DOCX Synthetic Pipeline tests PASSED!")

def test_pdf_pipeline():
    print_section("2. Testing PDF Pipeline & Direct Manipulation")
    
    tmp_dir = tempfile.mkdtemp(prefix="test_pdf_")
    test_pdf_path = os.path.join(tmp_dir, "sample_test.pdf")
    
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text(fitz.Point(72, 100), "RapidDoc PDF Evaluation Document", fontsize=16)
    page.insert_text(fitz.Point(72, 140), "This is a test paragraph on page one for verification.", fontsize=12)
    page.insert_text(fitz.Point(72, 180), "Status: Pending Review and Approval.", fontsize=12)
    pdf_doc.save(test_pdf_path)
    pdf_doc.close()
    print(f"Created temporary test PDF: {test_pdf_path}")
    
    # Test PDF Content Extraction
    content = get_pdf_content(test_pdf_path)
    print(f"Extracted {len(content)} pages. Page 1 blocks:")
    for block in content[0]["blocks"]:
        print(f"  [Block {block['block_no']}] {block['text']}")
    assert len(content[0]["blocks"]) >= 1, "Failed to extract PDF blocks"
    
    # Test Header & Footer Injection
    success = apply_pdf_styling(
        test_pdf_path, test_pdf_path,
        header_text="RapidDoc PDF Header",
        footer_text="RapidDoc PDF Footer - 2026"
    )
    print(f"Apply PDF Header/Footer result: {success}")
    assert success, "Failed to apply PDF header/footer"
    
    # Test Find & Replace in PDF
    rep_count = find_replace_pdf(
        test_pdf_path, test_pdf_path,
        find_text="Pending Review",
        replace_text="Approved & Verified"
    )
    print(f"PDF Find & Replace count: {rep_count}")
    
    shutil.rmtree(tmp_dir, ignore_errors=True)
    print("✅ PDF Synthetic Pipeline tests PASSED!")

def test_real_workspace_documents():
    print_section("3. Testing Real Documents in Workspace")
    
    # 1. Test real DOCX file: WR_1 (1).docx
    real_docx = r"c:\Users\umran\Desktop\RapidDoc_SGP\WR_1 (1).docx"
    if os.path.exists(real_docx):
        print(f"Testing real workspace DOCX: {real_docx}")
        content = get_docx_content(real_docx)
        img_count = get_docx_images_count(real_docx)
        hf = get_docx_headers_footers(real_docx)
        print(f"  -> Total paragraphs extracted: {len(content)}")
        print(f"  -> Total images found: {img_count}")
        print(f"  -> Headers/Footers: {hf}")
        print(f"  -> First 2 non-empty paragraphs:")
        non_empty = [p['text'] for p in content if p['text'].strip()][:2]
        for p in non_empty:
            print(f"     * {p[:80]}...")
    
    # 2. Test real PDF file: PROJECT PROPOSAL - RapidDoc.pdf
    real_pdf = r"c:\Users\umran\Desktop\RapidDoc_SGP\PROJECT PROPOSAL - RapidDoc.pdf"
    if os.path.exists(real_pdf):
        print(f"\nTesting real workspace PDF: {real_pdf}")
        pdf_content = get_pdf_content(real_pdf)
        pdf_img_count = get_pdf_images_count(real_pdf)
        pdf_hf = get_pdf_headers_footers(real_pdf)
        print(f"  -> Total pages extracted: {len(pdf_content)}")
        print(f"  -> Total images found: {pdf_img_count}")
        print(f"  -> Headers/Footers detected: {pdf_hf}")
        if pdf_content and pdf_content[0]["blocks"]:
            print(f"  -> First page block preview: {pdf_content[0]['blocks'][0]['text'][:80]}...")
            
    print("✅ Real Workspace Documents tests PASSED!")

def test_error_handling():
    print_section("4. Testing Comprehensive Error Handling & Edge Cases")
    
    # 1. Invalid file path handling
    print("Test 1: Non-existent file operations...")
    res_docx = get_docx_content("non_existent_file_path_12345.docx")
    print(f"  -> get_docx_content on non-existent: {res_docx}")
    assert res_docx == []
    
    res_pdf = get_pdf_content("non_existent_file_path_12345.pdf")
    print(f"  -> get_pdf_content on non-existent: {res_pdf}")
    assert res_pdf == []

    # 2. Corrupt / Empty document handling
    print("Test 2: Corrupted DOCX handling...")
    tmp_dir = tempfile.mkdtemp(prefix="test_err_")
    corrupt_docx = os.path.join(tmp_dir, "corrupt.docx")
    with open(corrupt_docx, "wb") as f:
        f.write(b"NOT_A_VALID_DOCX_STREAM_DATA")
    
    res = get_docx_content(corrupt_docx)
    print(f"  -> get_docx_content on corrupt returned: {res}")
    assert res == []
        
    # 3. Out-of-bounds DOCX update
    print("Test 3: Out-of-bounds index update...")
    clean_docx = os.path.join(tmp_dir, "clean.docx")
    d = docx.Document()
    d.add_paragraph("Only paragraph")
    d.save(clean_docx)
    
    res = update_docx_content(clean_docx, clean_docx, [{"index": 999, "text": "Out of bounds"}])
    print(f"  -> Out-of-bounds index update handled gracefully: result = {res}")
    assert res is True

    # 4. Empty and gibberish natural language commands
    print("Test 4: Empty & gibberish AI commands...")
    for cmd in ["", "   ", "xyzabc random gibberish 12345"]:
        understood = understand_command(cmd)
        print(f"  -> Command '{cmd}' -> Result: {understood}")
        assert "action" in understood
        
    # 5. Empty text rewrite
    print("Test 5: Empty rewrite text...")
    rw_res = rewrite_text("make formal", "")
    print(f"  -> Rewrite empty string: {rw_res}")
    assert rw_res["engine"] == "none"

    shutil.rmtree(tmp_dir, ignore_errors=True)
    print("✅ Error handling & edge case tests PASSED!")

if __name__ == "__main__":
    print("🚀 Starting RapidDoc Full Integration & Resilience Test Suite")
    test_docx_pipeline()
    test_pdf_pipeline()
    test_real_workspace_documents()
    test_error_handling()
    print("\n🎉 ALL PIPELINE & RESILIENCE TESTS COMPLETED SUCCESSFULLY!")
