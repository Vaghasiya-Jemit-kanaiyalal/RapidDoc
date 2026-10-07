import os
import io
import pytest
import docx
from docx.shared import Inches, Pt, RGBColor
from RapidDoc.backend.app.services.docx_editor import (
    update_docx_content,
    insert_docx_table,
    add_docx_table_row,
    add_docx_table_column,
    delete_docx_table_row,
    delete_docx_table_column,
    delete_docx_table,
    insert_docx_image,
    move_docx_image,
    update_docx_page_setup,
    insert_docx_page_break,
    apply_hierarchical_heading_numbers,
    style_docx_headings,
)

@pytest.fixture
def sample_docx(tmp_path):
    path = str(tmp_path / "sample.docx")
    doc = docx.Document()
    doc.add_paragraph("First paragraph text")
    doc.add_heading("Main Section", level=1)
    doc.add_paragraph("Second paragraph text under heading")
    doc.add_heading("Sub Section", level=2)
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "H1"
    t.rows[0].cells[1].text = "H2"
    t.rows[0].cells[2].text = "H3"
    t.rows[1].cells[0].text = "D1"
    t.rows[1].cells[1].text = "D2"
    t.rows[1].cells[2].text = "D3"
    doc.save(path)
    return path


def test_update_docx_content_with_runs_and_formatting(sample_docx, tmp_path):
    out_path = str(tmp_path / "formatted.docx")
    # Apply run-level bold to only one word
    runs = [
        {"text": "The ", "bold": False},
        {"text": "RapidDoc", "bold": True, "color": "2563eb"},
        {"text": " system is powerful.", "bold": False, "italic": True}
    ]
    edits = [{
        "index": 0,
        "runs": runs,
        "alignment": "center",
        "heading_level": 1
    }]
    success = update_docx_content(sample_docx, out_path, edits)
    assert success is True

    # Validate output docx
    reloaded = docx.Document(out_path)
    p0 = reloaded.paragraphs[0]
    assert p0.text == "The RapidDoc system is powerful."
    assert len(p0.runs) == 3
    assert p0.runs[0].bold is not True
    assert p0.runs[1].bold is True
    assert p0.runs[1].text == "RapidDoc"
    assert p0.runs[2].italic is True


def test_table_operations(sample_docx, tmp_path):
    out_path = str(tmp_path / "table_ops.docx")

    # Add row
    r1 = add_docx_table_row(sample_docx, out_path, table_index=0, position="below")
    assert r1["success"] is True

    # Add column
    r2 = add_docx_table_column(out_path, out_path, table_index=0, position="right", header_title="Status")
    assert r2["success"] is True

    # Delete column
    r3 = delete_docx_table_column(out_path, out_path, table_index=0, col_index=1)
    assert r3["success"] is True

    # Insert brand new table
    r4 = insert_docx_table(out_path, out_path, rows=4, cols=4, after_paragraph_index=0)
    assert r4["success"] is True

    reloaded = docx.Document(out_path)
    assert len(reloaded.tables) == 2


def test_page_setup_and_break(sample_docx, tmp_path):
    out_path = str(tmp_path / "page_ops.docx")

    # Landscape & Margins
    r1 = update_docx_page_setup(sample_docx, out_path, orientation="landscape", margin_inches=0.75, page_size="A4")
    assert r1["success"] is True

    # Page Break
    r2 = insert_docx_page_break(out_path, out_path, paragraph_index=0)
    assert r2["success"] is True

    reloaded = docx.Document(out_path)
    sec = reloaded.sections[0]
    assert sec.page_width > sec.page_height  # Landscape


def test_heading_numbering(sample_docx, tmp_path):
    out_path = str(tmp_path / "headings.docx")
    res = apply_hierarchical_heading_numbers(sample_docx, out_path)
    assert res["success"] is True

    reloaded = docx.Document(out_path)
    headings = [p.text for p in reloaded.paragraphs if p.style and p.style.name.lower().startswith("heading")]
    assert any("1. " in h for h in headings)
    assert any("1.1. " in h for h in headings)
