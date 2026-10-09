import pytest
from RapidDoc.backend.app.services.gemini_service import understand_command
from RapidDoc.backend.app.services.docx_editor import resize_docx_images, replace_docx_images, SIZE_MODE_ORIGINAL
import docx
import io
from PIL import Image

def test_ahange_all_header_to_blue():
    # User prompt: "ahange all header to blue" (handles typo ahange, header -> heading styling, blue color)
    res = understand_command("ahange all header to blue")
    assert res.get("action") == "style_headings"
    assert res.get("color") == "blue"
    assert res.get("color_rgb") == (0, 102, 204)

def test_heading_fonts_and_sizes():
    res1 = understand_command("change all headers to blue and font size 24")
    assert res1.get("action") == "style_headings"
    assert res1.get("color") == "blue"
    assert res1.get("font_size") == 24.0

    res2 = understand_command("make headings Arial and bold")
    assert res2.get("action") == "style_headings"
    assert res2.get("font_name") == "Arial"
    assert res2.get("bold") is True

    res3 = understand_command("set heading 1 size to 18pt")
    assert res3.get("action") == "style_headings"
    assert res3.get("font_size") == 18.0
    assert res3.get("level") == 1

def test_document_font_and_spacing():
    res = understand_command("change document font to Calibri and line spacing 1.5")
    assert res.get("action") == "style_document"
    assert res.get("font_name") == "Calibri"
    assert res.get("line_spacing") == 1.5

def test_image_resize_intents():
    res1 = understand_command("resize image 1 to 4 inches")
    assert res1.get("action") == "resize_image"
    assert res1.get("image_index") == 0
    assert res1.get("width") == 4.0
    assert res1.get("unit") == "in"

    res2 = understand_command("resize second image width 350px")
    assert res2.get("action") == "resize_image"
    assert res2.get("image_index") == 1
    assert res2.get("width") == 350.0
    assert res2.get("unit") == "px"

    res3 = understand_command("move image 1 to new page")
    assert res3.get("action") == "resize_image"
    assert res3.get("image_index") == 0
    assert res3.get("new_page") is True

def test_docx_image_resize_with_new_page(tmp_path):
    # Create a small docx with an image
    doc = docx.Document()
    doc.add_paragraph("First paragraph before image")

    img_stream = io.BytesIO()
    img = Image.new("RGB", (200, 200), color="red")
    img.save(img_stream, format="PNG")
    img_stream.seek(0)

    doc.add_picture(img_stream, width=docx.shared.Inches(2.0))
    doc_path = str(tmp_path / "test_doc.docx")
    out_path = str(tmp_path / "out_doc.docx")
    doc.save(doc_path)

    # Resize with new_page=True
    res = resize_docx_images(doc_path, out_path, [{
        "target_index": 0,
        "width": 3.0,
        "unit": "in",
        "new_page": True
    }])
    assert res.get("ok") is True

    # Verify that the image paragraph received <w:pageBreakBefore/>
    reloaded = docx.Document(out_path)
    found_break = False
    for p in reloaded.paragraphs:
        if p._element.xpath(".//w:pageBreakBefore"):
            found_break = True
            break
    assert found_break is True
