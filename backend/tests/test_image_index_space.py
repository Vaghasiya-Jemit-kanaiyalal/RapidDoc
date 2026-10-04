"""The image index space must be identical for reading and writing.

Every bug in this area is the same shape: the reader and the writer numbered
images differently, so a perfectly valid index from the UI replaced the wrong
picture - or errored. These tests pin one index space across all four consumers
(count, byte-serving, content extraction, replacement).
"""

import zipfile

from RapidDoc.backend.app.services.docx_editor import (
    collect_docx_image_targets,
    get_docx_image_parts,
    replace_docx_images,
)
from RapidDoc.backend.app.services.export_service import get_document_image_bytes
from RapidDoc.backend.app.services.pdf_editor import (
    get_pdf_content,
    get_pdf_image_slots,
    get_pdf_images_count,
    replace_pdf_images,
)

from imagefixtures import image_size, jpeg_bytes, pixel_at, png_bytes


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_shared_part_is_one_index_with_occurrences(docx_images):
    parts = get_docx_image_parts(docx_images)
    assert len(parts) == 3, "3 stored images, one of them drawn twice"
    assert [p["occurrences"] for p in parts] == [1, 2, 1]
    # Widths are the *displayed* size (EMU -> inches -> 96dpi px), not the stored
    # pixel count: the 500px figure is placed 5in wide = 480px, and the 40px logo
    # is placed at its native 72dpi size, which renders 53px at 96dpi. That is the
    # size the user sees in the editor, which is the number they are shown.
    assert [p["width"] for p in parts] == [480, 53, 768]


def test_docx_inline_shapes_index_space_differs_from_canonical(docx_images):
    """The reason the old writer broke: it used doc.inline_shapes.

    The counts here are identical (3 and 3), so a count-only test would pass
    while the numbering was completely different. Index 0 is the floating
    picture under the canonical walk and the logo under inline_shapes - so a UI
    built on inline_shapes would hand the writer an index that silently replaced
    the wrong image.
    """
    import docx

    d = docx.Document(docx_images)
    inline = d.inline_shapes
    canonical = collect_docx_image_targets(d)

    assert len(inline) == len(canonical) == 3
    inline_first_rid = inline[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
    assert canonical[0]["r_id"] != inline_first_rid, (
        "index 0 must be the floating picture, not the logo"
    )


def test_docx_index_space_survives_replacement(docx_images):
    before = get_docx_image_parts(docx_images)
    outcome = replace_docx_images(
        docx_images, docx_images, [{"target_index": 1, "image_bytes": png_bytes(64, 64)}]
    )
    assert outcome["ok"], outcome
    after = get_docx_image_parts(docx_images)
    assert [p["index"] for p in after] == [p["index"] for p in before]
    assert len(after) == len(before), "replacement must not renumber anything"


def test_docx_replacing_shared_part_updates_every_placement(docx_images):
    replace_docx_images(
        docx_images, docx_images, [{"target_index": 1, "image_bytes": png_bytes(64, 64)}]
    )
    with zipfile.ZipFile(docx_images) as z:
        document_xml = z.read("word/document.xml").decode("utf8")
    assert document_xml.count("<a:blip") == 4, "3 pictures + 1 shared extra draw"
    # One new media part, the old one left in place rather than overwritten:
    # overwriting the blob is what used to leave a PNG part holding JPEG bytes.
    media = [n for n in zipfile.ZipFile(docx_images).namelist() if n.startswith("word/media/")]
    assert len(media) == 4, media


def test_docx_cross_format_replacement_updates_content_type(docx_images):
    """Swapping PNG bytes for JPEG must not leave ``content_type: image/png``."""
    replace_docx_images(
        docx_images, docx_images, [{"target_index": 0, "image_bytes": jpeg_bytes(300, 300)}]
    )
    data, mime = get_document_image_bytes(docx_images, "docx", 0)
    assert mime == "image/jpeg", mime
    assert image_size(data) == (300, 300, "JPEG")


def test_docx_out_of_range_index_is_a_clear_error(docx_images):
    outcome = replace_docx_images(
        docx_images, docx_images, [{"target_index": 99, "image_bytes": png_bytes(10, 10)}]
    )
    assert not outcome["ok"]
    assert "does not exist" in outcome["error"]


def test_docx_bad_index_does_not_block_the_good_one(docx_images):
    outcome = replace_docx_images(docx_images, docx_images, [
        {"target_index": 99, "image_bytes": png_bytes(10, 10)},
        {"target_index": 0, "image_bytes": png_bytes(70, 70)},
    ])
    assert outcome["replaced"] == 1


def test_docx_missing_extent_does_not_crash(docx_images):
    """No declared size anywhere must still swap, and say so rather than guess."""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(docx_images)
    target = d.element.body.findall(".//" + qn("w:drawing"))[0]
    for tag in ("wp:extent", "a:ext"):
        for node in target.findall(".//" + qn(tag)):
            node.getparent().remove(node)
    d.save(docx_images)

    outcome = replace_docx_images(
        docx_images, docx_images, [{"target_index": 0, "image_bytes": png_bytes(90, 30)}]
    )
    assert outcome["ok"], outcome
    assert outcome["reports"][0]["no_extent"] is True


def test_docx_falls_back_to_xfrm_extent_when_wp_extent_is_gone(docx_images):
    """a:ext inside a:xfrm is a valid size source; it must be used, not ignored."""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(docx_images)
    target = d.element.body.findall(".//" + qn("w:drawing"))[0]
    for node in target.findall(".//" + qn("wp:extent")):
        node.getparent().remove(node)
    d.save(docx_images)

    outcome = replace_docx_images(
        docx_images, docx_images, [{"target_index": 0, "image_bytes": png_bytes(90, 30)}]
    )
    assert outcome["ok"], outcome
    report = outcome["reports"][0]
    assert "no_extent" not in report, "a:ext was available, so this is not an extent-less drawing"
    assert report["scaled"] is True, "the aspect ratio was fitted using the a:ext size"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_indexes_occurrences_not_distinct_xrefs(pdf_images):
    slots = get_pdf_image_slots(pdf_images)
    assert len(slots) == 5
    assert get_pdf_images_count(pdf_images) == len(slots)
    assert [s["page_num"] for s in slots] == [0, 0, 1, 2, 2]


def test_pdf_content_uses_the_same_index_space(pdf_images):
    slots = get_pdf_image_slots(pdf_images)
    pages = get_pdf_content(pdf_images)
    seen = [img["image_index"] for page in pages for img in page["images"]]
    assert seen == [s["index"] for s in slots]


def test_pdf_slots_carry_placement_geometry(pdf_images):
    """Reading-order numbering depends on it; resizing depends on it."""
    for slot in get_pdf_image_slots(pdf_images):
        assert slot["bbox"] is not None
        x0, y0, x1, y1 = slot["bbox"]
        assert x1 > x0 and y1 > y0


def test_pdf_repeated_xref_reports_all_occurrences(pdf_repeated_logo):
    slots = get_pdf_image_slots(pdf_repeated_logo)
    assert len(slots) == 3
    assert {s["xref"] for s in slots} == {slots[0]["xref"]}
    assert {s["occurrences"] for s in slots} == {3}


def test_pdf_replacement_keeps_slot_count_and_position(pdf_images):
    before = get_pdf_image_slots(pdf_images)
    outcome = replace_pdf_images(
        pdf_images, pdf_images, [{"target_index": 2, "image_bytes": png_bytes(60, 60)}]
    )
    assert outcome["ok"], outcome
    after = get_pdf_image_slots(pdf_images)
    assert len(after) == len(before), "must not add a duplicate resource"
    # Same drawn rectangle: only the pixels inside it changed.
    for b, a in zip(before, after):
        assert [round(v, 2) for v in b["bbox"]] == [round(v, 2) for v in a["bbox"]]


def test_pdf_replacement_changes_pixels(pdf_images):
    """Index 1 is the 2:1 wide box; a 3:1 replacement is letterboxed, so the
    centre is the only reliable place to sample."""
    yellow = png_bytes(300, 100, (255, 255, 0))
    replace_pdf_images(
        pdf_images, pdf_images, [{"target_index": 1, "image_bytes": yellow}]
    )
    data, _ = get_document_image_bytes(pdf_images, "pdf", 1)
    assert pixel_at(data, 150, 50) == (255, 255, 0)
    assert pixel_at(data, 2, 2) == (255, 255, 255), "letterbox padding stays white"


def test_pdf_replacing_one_shared_slot_reports_the_others(pdf_repeated_logo):
    outcome = replace_pdf_images(
        pdf_repeated_logo, pdf_repeated_logo,
        [{"target_index": 1, "image_bytes": png_bytes(40, 40, (7, 7, 7))}],
    )
    assert outcome["ok"], outcome
    assert outcome["reports"][0]["occurrences"] == 3, "user must be told all 3 changed"


def test_pdf_out_of_range_index_is_a_clear_error(pdf_images):
    outcome = replace_pdf_images(
        pdf_images, pdf_images, [{"target_index": 99, "image_bytes": png_bytes(10, 10)}]
    )
    assert not outcome["ok"]
    assert "does not exist" in outcome["error"]


def test_document_with_no_images_reports_zero(docx_no_images):
    assert get_docx_image_parts(docx_no_images) == []