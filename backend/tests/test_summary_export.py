"""Summary exports: txt, docx and pdf.

The export is judged on what the user ends up holding, so the assertions read the
produced file back with the same libraries that will open it - text is checked as
text, DOCX through python-docx, PDF through PyMuPDF. A summary that "exports
fine" but produces an empty or garbled file is the failure mode worth catching.

The PDF path shells out to LibreOffice, which is genuinely optional on a machine,
so that one test skips rather than failing when the converter is unavailable.
"""

import io

import docx
import fitz
import pytest
from bson import ObjectId
from fastapi import FastAPI
from fastapi.testclient import TestClient

USER = {"id": "user-1", "email": "user@example.com", "role": "user"}
DOC_ID = "0" * 24
BASE = f"/api/documents/{DOC_ID}"

PAYLOAD = {
    "summary": "The bridge project replaced two obsolete spans and added a "
               "cable-stayed main span over the shipping channel.",
    "key_points": [
        "Main span is cable-stayed",
        "Two obsolete spans removed",
        "Budget came in under estimate",
    ],
    "source": "whole document",
    "engine": "local",
    "characters": 18422,
}


@pytest.fixture
def client(tmp_path, monkeypatch):
    from RapidDoc.backend.app.routers import auth, documents

    doc = {
        "_id": ObjectId(DOC_ID),
        "owner_id": "user-1",
        "file_type": "pdf",
        "filename": "Bridge Report.pdf",
        "storage_path": "bridge.pdf",
        "edit_history": [],
    }

    class FakeCollection:
        def __init__(self, document):
            self.doc = document

        def find_one(self, query=None, *args, **kwargs):
            return self.doc

        def update_one(self, query, update):
            for entry in (update.get("$push") or {}).values():
                self.doc.setdefault("edit_history", []).extend(
                    entry if isinstance(entry, list) else [entry]
                )
            self.doc.update(update.get("$set") or {})
            return None

    class FakeDb:
        def __init__(self, document):
            self.documents = FakeCollection(document)

    db = FakeDb(doc)
    monkeypatch.setattr(documents.db_conn, "get_db", lambda: db)

    app = FastAPI()
    app.include_router(documents.router)
    app.dependency_overrides[auth.get_current_user] = lambda: USER
    test_client = TestClient(app)
    test_client.doc = doc
    return test_client


def export(client, fmt, **overrides):
    body = {"format": fmt, **PAYLOAD, **overrides}
    return client.post(f"{BASE}/summary/export", json=body)


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------

def test_txt_export_carries_the_summary_and_the_points(client):
    response = export(client, "txt")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/plain")

    text = response.content.decode("utf-8")
    assert PAYLOAD["summary"] in text
    assert "Main span is cable-stayed" in text
    # Provenance is the part that makes an exported summary auditable later.
    assert "whole document" in text
    assert "local" in text


def test_txt_export_is_named_after_the_document(client):
    response = export(client, "txt")
    disposition = response.headers["content-disposition"]
    assert "Bridge_Report.txt" in disposition


def test_an_explicit_title_overrides_the_document_name(client):
    response = export(client, "txt", title="Executive summary")
    assert "Executive_summary.txt" in response.headers["content-disposition"]
    assert "Executive summary" in response.content.decode("utf-8")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_export_uses_real_list_items(client):
    """Bullets have to be list items, not paragraphs starting with a dash."""
    response = export(client, "docx")
    assert response.status_code == 200, response.text

    document = docx.Document(io.BytesIO(response.content))
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullets == PAYLOAD["key_points"]

    body = "\n".join(p.text for p in document.paragraphs)
    assert PAYLOAD["summary"] in body


def test_docx_export_is_a_valid_office_file(client):
    response = export(client, "docx")
    assert response.content[:2] == b"PK"  # zip container, i.e. a real .docx
    assert "wordprocessingml" in response.headers["content-type"]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_export_produces_readable_pages(client):
    response = export(client, "pdf")
    if response.status_code == 503:
        pytest.skip("LibreOffice is unavailable, so DOCX -> PDF cannot run here.")

    assert response.status_code == 200, response.text
    assert response.content[:5] == b"%PDF-"

    with fitz.open(stream=response.content, filetype="pdf") as pdf:
        assert pdf.page_count >= 1
        rendered = "".join(page.get_text() for page in pdf)
    assert "Main span is cable-stayed" in rendered


# ---------------------------------------------------------------------------
# Failure modes
# ---------------------------------------------------------------------------

def test_an_empty_summary_is_rejected(client):
    response = export(client, "txt", summary="   ")
    assert response.status_code == 400
    assert "no summary text" in response.json()["detail"].lower()


def test_an_unsupported_format_is_rejected_with_the_options(client):
    response = export(client, "rtf")
    assert response.status_code == 400
    assert "txt" in response.json()["detail"]
    assert "docx" in response.json()["detail"]


def test_a_missing_summary_is_a_pydantic_422(client):
    response = client.post(f"{BASE}/summary/export", json={"format": "txt"})
    assert response.status_code == 422


def test_the_export_is_recorded_in_the_edit_log(client):
    export(client, "txt")
    actions = [entry["action"] for entry in client.doc["edit_history"]]
    assert "Exported summary as TXT" in actions


def test_duplicate_points_are_collapsed(client):
    """A model returning the same bullet twice must not pad a one-page summary."""
    from RapidDoc.backend.app.services.summary_export import build_summary_txt

    payload = {"summary": "S", "key_points": ["Same thing.", "same thing", "Other"]}
    text = build_summary_txt(payload, "T").decode("utf-8")
    assert text.count("Same thing.") == 1
    assert text.count("Other") == 1


def test_a_summary_without_key_points_still_exports(client):
    response = export(client, "txt", key_points=[])
    assert response.status_code == 200
    assert PAYLOAD["summary"] in response.content.decode("utf-8")